import io
import pytest
from fastapi.testclient import TestClient
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool
from docx import Document as DocxDocument

from backend.app.main import app
from backend.app.core.database import Base, get_db

client = TestClient(app)


@pytest.fixture
def auth_headers():
    # Register and login test user
    client.post("/api/auth/register", json={
        "email": "doc_tester@medassist.ai",
        "password": "Password123!"
    })
    login_res = client.post("/api/auth/login", json={
        "email": "doc_tester@medassist.ai",
        "password": "Password123!"
    })
    token = login_res.json()["access_token"]
    return {"Authorization": f"Bearer {token}"}


def test_upload_and_process_txt_document(auth_headers):
    sample_text = (
        "CLINICAL STUDY ON HYPERTENSION\n\n"
        "Abstract:\n"
        "Essential hypertension affects over one billion individuals worldwide. "
        "Primary risk factors include high sodium intake, sedentary lifestyle, and genetic predispositions.\n\n"
        "Methodology:\n"
        "A cohort of 500 adult patients participated in an educational lifestyle modification trial over 12 weeks.\n\n"
        "Laboratory Findings:\n"
        "Mean resting systolic blood pressure decreased by 8.4 mm Hg in the intervention cohort."
    )
    files = {"file": ("hypertension_study.txt", io.BytesIO(sample_text.encode("utf-8")), "text/plain")}

    response = client.post("/api/documents/upload", headers=auth_headers, files=files)
    assert response.status_code == 201
    data = response.json()
    assert data["filename"] == "hypertension_study.txt"
    assert data["status"] == "COMPLETED"
    assert data["file_type"] == "txt"
    doc_id = data["id"]

    # Check status endpoint
    status_res = client.get(f"/api/documents/{doc_id}/status", headers=auth_headers)
    assert status_res.status_code == 200
    status_data = status_res.json()
    assert status_data["status"] == "COMPLETED"
    assert status_data["progress_percent"] == 100
    assert status_data["chunk_count"] > 0

    # Check details endpoint
    detail_res = client.get(f"/api/documents/{doc_id}", headers=auth_headers)
    assert detail_res.status_code == 200
    detail_data = detail_res.json()
    assert detail_data["title"] is not None
    assert len(detail_data["chunks"]) > 0

    # Check list documents
    list_res = client.get("/api/documents", headers=auth_headers)
    assert list_res.status_code == 200
    assert len(list_res.json()) >= 1


def test_upload_and_process_csv_document(auth_headers):
    csv_content = (
        "PatientID,Age,Gender,SystolicBP,DiastolicBP,Cholesterol\n"
        "P1001,54,Male,138,88,210\n"
        "P1002,62,Female,142,90,230\n"
        "P1003,45,Male,122,80,185\n"
    )
    files = {"file": ("lab_metrics.csv", io.BytesIO(csv_content.encode("utf-8")), "text/csv")}

    response = client.post("/api/documents/upload", headers=auth_headers, files=files)
    assert response.status_code == 201
    data = response.json()
    assert data["status"] == "COMPLETED"
    assert data["category"] == "laboratory_report"


def test_upload_and_process_docx_document(auth_headers):
    # Create docx in memory
    doc = DocxDocument()
    doc.add_heading("Cardiovascular Risk Assessment", level=1)
    doc.add_paragraph("Cardiovascular diseases remain the leading cause of mortality globally.")
    doc.add_heading("Recommendations", level=2)
    doc.add_paragraph("Routine aerobic exercise and DASH dietary patterns demonstrate positive outcomes.")

    file_stream = io.BytesIO()
    doc.save(file_stream)
    file_stream.seek(0)

    files = {"file": ("cardio_risk.docx", file_stream, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
    response = client.post("/api/documents/upload", headers=auth_headers, files=files)
    assert response.status_code == 201
    assert response.json()["status"] == "COMPLETED"


def test_reject_unsupported_file_extension(auth_headers):
    files = {"file": ("script.exe", io.BytesIO(b"malicious content"), "application/octet-stream")}
    response = client.post("/api/documents/upload", headers=auth_headers, files=files)
    assert response.status_code == 400
    assert "Unsupported file type" in response.json()["detail"]


def test_reject_fake_pdf_header(auth_headers):
    # File named .pdf but lacking %PDF- magic bytes
    files = {"file": ("corrupt.pdf", io.BytesIO(b"Not a real PDF file content"), "application/pdf")}
    response = client.post("/api/documents/upload", headers=auth_headers, files=files)
    assert response.status_code == 400
    assert "Invalid PDF format" in response.json()["detail"]


def test_delete_document_and_cascade(auth_headers):
    files = {"file": ("to_delete.txt", io.BytesIO(b"Temporary clinical notes."), "text/plain")}
    res = client.post("/api/documents/upload", headers=auth_headers, files=files)
    doc_id = res.json()["id"]

    del_res = client.delete(f"/api/documents/{doc_id}", headers=auth_headers)
    assert del_res.status_code == 200

    # Ensure 404 on subsequent get
    get_res = client.get(f"/api/documents/{doc_id}", headers=auth_headers)
    assert get_res.status_code == 404
